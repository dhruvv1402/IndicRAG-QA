"""Render paper/slides.md into a .pptx deck.

The markdown is the source. It is easier to edit, it diffs, and the speaker
notes sit next to the slide they belong to; the .pptx is a build artefact, the
same as paper.md and the .docx.

Speaker notes are carried into the notes pane rather than dropped. The notes in
slides.md are where the argument lives -- which row to read aloud, where to
pause, what to cut if time runs short -- and a deck that loses them is a deck of
bullet points.

    python scripts/build-pptx.py

Writes paper/IndicRAG-QA-slides.pptx.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "paper" / "slides.md"
TARGET = ROOT / "paper" / "IndicRAG-QA-slides.pptx"

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
ACCENT = RGBColor(0x2C, 0x5F, 0x8A)
BAD = RGBColor(0xB0, 0x4A, 0x4A)

FONT = "Calibri"


def strip_md(text: str) -> str:
    """Slide text is plain; emphasis is carried by layout, not by asterisks."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"^>\s*", "", text)
    return text.strip()


def parse(md: str) -> list[dict]:
    """Split slides.md on `## ` headings into slide records."""
    slides: list[dict] = []
    current: dict | None = None
    in_notes = False

    for line in md.splitlines():
        stripped = line.strip()

        if stripped.startswith("## "):
            if current:
                slides.append(current)
            title = re.sub(r"^\d+\.\s*", "", stripped[3:]).strip()
            current = {"title": title, "body": [], "notes": [], "table": []}
            in_notes = False
            continue

        if current is None:
            continue  # deck preamble

        if stripped.startswith("*Notes:*"):
            in_notes = True
            current["notes"].append(strip_md(stripped[8:]))
            continue

        if re.fullmatch(r"-{3,}", stripped):
            in_notes = False
            continue

        if not stripped:
            continue

        if in_notes:
            current["notes"].append(strip_md(stripped))
        elif stripped.startswith("|"):
            raw = stripped.strip("|")
            if not re.fullmatch(r"[\s:|-]+", raw):
                current["table"].append([strip_md(c) for c in raw.split("|")])
        else:
            current["body"].append(strip_md(stripped))

    if current:
        slides.append(current)
    return slides


def add_textbox(slide, left, top, width, height, lines, *, size, bold=False, colour=INK):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = para.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = FONT
        run.font.color.rgb = colour
        para.space_after = Pt(6)
    return box


def add_table(slide, rows, left, top, width):
    n_rows, n_cols = len(rows), len(rows[0])
    height = Inches(0.32) * n_rows
    shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = shape.table

    for r, row in enumerate(rows):
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.text = row[c] if c < len(row) else ""
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(13)
                    run.font.name = FONT
                    run.font.bold = r == 0
                    run.font.color.rgb = INK
    return shape


def build() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    slides = parse(SOURCE.read_text(encoding="utf-8"))

    for n, rec in enumerate(slides):
        slide = prs.slides.add_slide(blank)

        is_title = n == 0
        # The deck's first section is headed "## 1. Title" -- a label for the
        # author, not the title itself, which sits in the body beneath it.
        # Promote it, or the opening slide reads "Title".
        if is_title and rec["body"]:
            rec["title"] = rec["body"].pop(0)

        add_textbox(
            slide, Inches(0.7), Inches(0.5) if not is_title else Inches(2.4),
            Inches(12.0), Inches(1.0), [rec["title"]],
            size=40 if is_title else 28, bold=True,
            colour=ACCENT if is_title else INK,
        )

        top = Inches(3.4) if is_title else Inches(1.6)

        if rec["body"]:
            add_textbox(
                slide, Inches(0.7), top, Inches(12.0), Inches(3.4), rec["body"],
                size=20 if is_title else 18,
                colour=MUTED if is_title else INK,
            )
            top = Emu(top + Inches(0.42) * len(rec["body"]))

        if rec["table"]:
            add_table(slide, rec["table"], Inches(0.7), top, Inches(11.0))

        if rec["notes"]:
            slide.notes_slide.notes_text_frame.text = " ".join(rec["notes"])

        # Slide number, bottom right.
        add_textbox(
            slide, Inches(12.3), Inches(6.9), Inches(0.7), Inches(0.35),
            [str(n + 1)], size=11, colour=MUTED,
        )

    return prs


def main() -> int:
    if not SOURCE.exists():
        print(f"missing {SOURCE}")
        return 1
    prs = build()
    prs.save(TARGET)
    notes = sum(
        1 for s in prs.slides
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()
    )
    print(f"wrote {TARGET.relative_to(ROOT)}  ({len(prs.slides._sldIdLst)} slides, "
          f"{notes} with speaker notes, {TARGET.stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
